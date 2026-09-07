import os
import re
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def convert_markdown_to_docx(md_path: Path, docx_path: Path):
    doc = docx.Document()

    # Set standard 1-inch margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Set base styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Malgun Gothic'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    normal_style.paragraph_format.line_spacing = 1.2
    normal_style.paragraph_format.space_after = Pt(4)

    def set_cell_background(cell, fill_hex):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tcPr.append(shd)

    def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
            node = OxmlElement(f'w:{m}')
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    def set_table_borders(table):
        tblPr = table._tbl.tblPr
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'  <w:top w:val="single" w:sz="8" w:space="0" w:color="4A5568"/>'
            f'  <w:bottom w:val="single" w:sz="8" w:space="0" w:color="4A5568"/>'
            f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="CBD5E0"/>'
            f'  <w:insideV w:val="none"/>'
            f'  <w:left w:val="none"/>'
            f'  <w:right w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr.append(borders)

    def add_formatted_text(paragraph, text):
        pattern = re.compile(r'(\*\*.*?\*\*|\*.*?\*|`.*?`|\{doi:.*?\})')
        tokens = pattern.split(text)
        for token in tokens:
            if not token:
                continue
            if token.startswith('**') and token.endswith('**'):
                run = paragraph.add_run(token[2:-2])
                run.bold = True
                run.font.name = 'Malgun Gothic'
            elif token.startswith('*') and token.endswith('*') and not token.startswith('**'):
                run = paragraph.add_run(token[1:-1])
                run.italic = True
                run.font.name = 'Malgun Gothic'
            elif token.startswith('`') and token.endswith('`'):
                run = paragraph.add_run(token[1:-1])
                run.font.name = 'Consolas'
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(0x80, 0x00, 0x20)
            elif token.startswith('{doi:') and token.endswith('}'):
                run = paragraph.add_run(token)
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(0x15, 0x65, 0xC0)
            else:
                run = paragraph.add_run(token)
                run.font.name = 'Malgun Gothic'

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_table = False
    table_lines = []

    def flush_table():
        nonlocal table_lines
        if not table_lines:
            return
        rows = []
        for tl in table_lines:
            if re.match(r'^\|[\s\-:|]+\|$', tl):
                continue  # separator row
            cells = [c.strip() for c in tl.strip('|').split('|')]
            rows.append(cells)
        
        if rows:
            num_cols = max(len(r) for r in rows)
            for r in rows:
                while len(r) < num_cols:
                    r.append('')
            
            tbl = doc.add_table(rows=len(rows), cols=num_cols)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            set_table_borders(tbl)
            
            for r_idx, row in enumerate(rows):
                for c_idx, cell_value in enumerate(row):
                    cell = tbl.cell(r_idx, c_idx)
                    set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
                    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                    
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.line_spacing = 1.15
                    
                    if r_idx == 0:
                        set_cell_background(cell, 'EDF2F7')
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    
                    parts = re.split(r'<br\s*/?>', cell_value, flags=re.IGNORECASE)
                    for p_idx, part in enumerate(parts):
                        if p_idx > 0:
                            p = cell.add_paragraph()
                            p.paragraph_format.space_after = Pt(2)
                            p.paragraph_format.line_spacing = 1.15
                        if r_idx == 0:
                            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        
                        add_formatted_text(p, part.strip())
                        if r_idx == 0:
                            for r in p.runs:
                                r.bold = True
                                r.font.size = Pt(9.5)
                        else:
                            for r in p.runs:
                                r.font.size = Pt(9.0)
            
            doc.add_paragraph()
        table_lines = []

    for line_raw in lines:
        line = line_raw.rstrip('\r\n')
        
        if line.strip().startswith('|') and line.strip().endswith('|'):
            in_table = True
            table_lines.append(line.strip())
            continue
        elif in_table:
            in_table = False
            flush_table()
        
        img_match = re.match(r'^!\[(.*?)\]\((.*?)\)$', line.strip())
        if img_match:
            alt_text, img_rel_path = img_match.groups()
            img_path = Path(img_rel_path)
            if img_path.exists():
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(8)
                p_img.paragraph_format.space_after = Pt(4)
                run = p_img.add_run()
                run.add_picture(str(img_path), width=Inches(6.0))
            continue
        
        caption_match = re.match(r'^\*(그림\s*\d+\..*?)\*$', line.strip())
        if caption_match:
            cap_text = caption_match.group(1)
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(2)
            p_cap.paragraph_format.space_after = Pt(10)
            run = p_cap.add_run(cap_text)
            run.italic = True
            run.font.size = Pt(9.0)
            run.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)
            continue
        
        if line.startswith('# '):
            title_text = line[2:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(14)
            run = p.add_run(title_text)
            run.bold = True
            run.font.size = Pt(18)
            run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
            continue
        
        if line.startswith('## '):
            h2_text = line[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(h2_text)
            run.bold = True
            run.font.size = Pt(13.5)
            run.font.color.rgb = RGBColor(0x00, 0x4D, 0x40)
            continue
        
        if line.startswith('### '):
            h3_text = line[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(h3_text)
            run.bold = True
            run.font.size = Pt(11.5)
            run.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
            continue
        
        if line.startswith('#### '):
            h4_text = line[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(h4_text)
            run.bold = True
            run.font.size = Pt(10.5)
            run.font.color.rgb = RGBColor(0x31, 0x97, 0x95)
            continue
        
        if line.strip() == '---':
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            continue
        
        if re.match(r'^\s*[-•]\s+', line):
            bullet_text = re.sub(r'^\s*[-•]\s+', '', line).strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.2
            add_formatted_text(p, bullet_text)
            continue
        
        if re.match(r'^\s*\d+\.\s+', line):
            num_text = line.strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.2
            add_formatted_text(p, num_text)
            continue
        
        if line.strip():
            if re.match(r'^\*\*표\s*\d+.*?\*\*$', line.strip()):
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.space_after = Pt(4)
                add_formatted_text(p, line.strip())
            else:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.line_spacing = 1.25
                add_formatted_text(p, line.strip())

    if in_table:
        flush_table()

    doc.save(str(docx_path))
    print(f"Successfully generated {docx_path} ({docx_path.stat().st_size:,} bytes)")

if __name__ == '__main__':
    md = Path('manuscript.md')
    docx_out = Path('manuscript.docx')
    convert_markdown_to_docx(md, docx_out)
